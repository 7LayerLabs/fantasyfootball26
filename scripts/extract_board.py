"""Pull computed values from the recalced workbook into board data JSON."""
import openpyxl, json, warnings
from pathlib import Path

warnings.filterwarnings("ignore")

SCRIPT_DIR = Path(__file__).parent
WORKSPACE = SCRIPT_DIR.parent
P = WORKSPACE / "2026_Fantasy_Football_Draft_Engine_100_PERCENT_FIXED.xlsx"
PLAYERS_JSON = SCRIPT_DIR / "players.json"

wb = openpyxl.load_workbook(P, data_only=True)
de, ls = wb["Draft Engine"], wb["LIVE SOURCE"]

def tier(r):
    for t, lo in [("ELITE", 95), ("TIER 1", 90), ("TIER 2", 85), ("TIER 3", 80),
                  ("TIER 4", 75), ("TIER 5", 70), ("TIER 6", 65)]:
        if r >= lo: return t
    return "TIER 7"

# Load opportunity data from players.json
opp_data = {}
try:
    with open(PLAYERS_JSON, encoding="utf-8") as f:
        players_json = json.load(f)
        for p in players_json["players"]:
            opp_data[p["player"]] = {
                "snap_share": p.get("snap_share", 0),
                "route_share": p.get("route_share", 0),
                "carry_share": p.get("carry_share", 0),
                "target_share": p.get("target_share", 0),
                "rz_share": p.get("rz_share", 0),
                "expected_ppg": p.get("expected_ppg", 0),
                "actual_ppg": p.get("actual_ppg", 0),
                "games_played": p.get("games_played", 0),
                "sample_note": p.get("sample_note", ""),
                "qb_situation": p.get("qb_situation", ""),
                "role": p.get("role", ""),
                "opp_score": p.get("opp_score", 0)
            }
except Exception as e:
    print(f"Warning: Could not load opportunity data from players.json: {e}")

out = []
for r in range(3, 233):
    rating = round(de.cell(r, 5).value, 1)
    name = de.cell(r, 2).value
    opp = opp_data.get(name, {})
    
    out.append({
        "pos": de.cell(r, 1).value,
        "name": name,
        "team": ls.cell(r, 3).value or "",
        "bye": ls.cell(r, 12).value or 0,
        "rating": rating,
        "tier": tier(rating),
        "posRank": de.cell(r, 6).value,
        "pts": round(de.cell(r, 7).value, 1),
        "vorp": round(de.cell(r, 10).value, 1),
        "score": round(de.cell(r, 14).value, 1),
        "overall": de.cell(r, 17).value,
        "adp": round(de.cell(r, 18).value, 1),
        "edge": round(de.cell(r, 19).value, 1),
        "conf": de.cell(r, 20).value,
        "floor": round(de.cell(r, 21).value, 1),
        "ceil": round(de.cell(r, 23).value, 1),
        "rookie": de.cell(r, 24).value == "ROOKIE ENGINE",
        "exp": ls.cell(r, 8).value or "",
        "risk": ls.cell(r, 10).value or "",
        # Add opportunity fields
        "snap_share": opp.get("snap_share", 0),
        "route_share": opp.get("route_share", 0),
        "carry_share": opp.get("carry_share", 0),
        "target_share": opp.get("target_share", 0),
        "rz_share": opp.get("rz_share", 0),
        "expected_ppg": opp.get("expected_ppg", 0),
        "actual_ppg": opp.get("actual_ppg", 0),
        "games_played": opp.get("games_played", 0),
        "sample_note": opp.get("sample_note", ""),
        "qb_situation": opp.get("qb_situation", ""),
        "role": opp.get("role", ""),
        "opp_score": opp.get("opp_score", 0)
    })
out.sort(key=lambda p: p["overall"])
assert len(out) == 230 and [p["overall"] for p in out] == list(range(1, 231))

with open(SCRIPT_DIR / "board_data.json", "w", encoding="utf-8") as f:
    json.dump(out, f)

print("OK", len(out), "players;", out[0]["name"], "->", out[-1]["name"])
if opp_data:
    print(f"Added opportunity data for {len([p for p in out if p['games_played'] > 0])} players with NFL data")
